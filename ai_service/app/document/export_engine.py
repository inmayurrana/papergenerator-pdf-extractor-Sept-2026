import re
import io
import logging
from pathlib import Path
from typing import Dict, Any, List, Optional
import pymupdf  # type: ignore
import docx  # type: ignore
from .office_converter import extract_paragraph_with_math, extract_cell_with_math, get_unicode_font_args
from ..engines.specialized_math import specialized_math

logger = logging.getLogger("export_engine")

class ExportEngine:
    """
    Engine for:
    1. Generating professional A4 PDFs for Question Bank and Question Papers with Answer Keys.
    2. Generating formatted PDFs for extracted and translated text.
    3. Parsing and extracting structured questions and answers from Word (.docx) and PDF (.pdf) files.
    """

    PAGE_WIDTH = 595   # Standard A4 width in points
    PAGE_HEIGHT = 842  # Standard A4 height in points
    MARGIN_LEFT = 45
    MARGIN_RIGHT = 550
    MARGIN_TOP = 50
    MARGIN_BOTTOM = 790
    CONTENT_WIDTH = MARGIN_RIGHT - MARGIN_LEFT

    @classmethod
    def _create_page(
        cls,
        doc: pymupdf.Document,
        title: str = "",
        page_num: int = 1,
        margin_left: Optional[int] = None,
        margin_right: Optional[int] = None,
        margin_bottom: Optional[int] = None,
    ) -> Any:
        """Helper to create a new page with standard footer."""
        page = doc.new_page(width=cls.PAGE_WIDTH, height=cls.PAGE_HEIGHT)
        ml = margin_left if margin_left is not None else cls.MARGIN_LEFT
        mr = margin_right if margin_right is not None else cls.MARGIN_RIGHT
        mb = margin_bottom if margin_bottom is not None else cls.MARGIN_BOTTOM
        # Footer
        footer_text = f"{title}  •  Page {page_num}" if title else f"Page {page_num}"
        font_args = get_unicode_font_args(False)
        page.insert_textbox(
            pymupdf.Rect(ml, mb + 10, mr, mb + 30),
            footer_text,
            fontsize=8,
            color=(0.4, 0.4, 0.4),
            **font_args
        )
        return page

    @classmethod
    def generate_questions_pdf(
        cls,
        questions: List[Dict[str, Any]],
        title: str = "Question Bank & Answer Key",
        folder_name: str = "General",
        include_answers: bool = True,
    ) -> bytes:
        """Generates an A4 PDF containing Questions & Options, followed by an Answer Key & Solutions Guide."""
        doc = pymupdf.open()
        page_num = 1
        page = cls._create_page(doc, title, page_num)
        y = cls.MARGIN_TOP

        # Title Header
        page.insert_text((cls.MARGIN_LEFT, y + 16), title.upper(), fontsize=15, fontname="hebo", color=(0.1, 0.15, 0.35))
        y += 24
        meta_str = f"Taxonomy Folder: {folder_name}  |  Total Questions: {len(questions)}  |  Exported with Solutions"
        page.insert_text((cls.MARGIN_LEFT, y + 10), meta_str, fontsize=9, fontname="helv", color=(0.3, 0.3, 0.3))
        y += 18
        page.draw_line((cls.MARGIN_LEFT, y), (cls.MARGIN_RIGHT, y), color=(0.2, 0.3, 0.6), width=1.5)
        y += 15

        # Section 1: Questions & Options
        page.insert_text((cls.MARGIN_LEFT, y + 10), "PART I: EXAMINATION QUESTIONS", fontsize=11, fontname="hebo", color=(0.1, 0.1, 0.2))
        y += 18

        for idx, q in enumerate(questions):
            # Check for page overflow
            if y > cls.MARGIN_BOTTOM - 80:
                page_num += 1
                page = cls._create_page(doc, title, page_num)
                y = cls.MARGIN_TOP

            q_num = q.get("questionNumber") or str(idx + 1)
            marks = q.get("marks", 1)
            q_text = q.get("questionText") or ""
            opts = q.get("options") or []
            if isinstance(opts, str):
                try:
                    import json
                    opts = json.loads(opts)
                except Exception:
                    opts = []

            # Question stem header with marks
            marks_str = f"[{marks} Mark{'s' if marks > 1 else ''}]"
            marks_w = pymupdf.get_text_length(marks_str, fontname="hebo", fontsize=9)
            page.insert_text((cls.MARGIN_RIGHT - marks_w, y + 10), marks_str, fontsize=9, fontname="hebo", color=(0.2, 0.4, 0.2))
            
            # Question stem
            stem_prefix = f"Q{q_num}. "
            prefix_w = pymupdf.get_text_length(stem_prefix, fontname="hebo", fontsize=10)
            page.insert_text((cls.MARGIN_LEFT, y + 10), stem_prefix, fontsize=10, fontname="hebo", color=(0.1, 0.1, 0.1))

            rect = pymupdf.Rect(cls.MARGIN_LEFT + prefix_w, y, cls.MARGIN_RIGHT - marks_w - 10, y + 120)
            rc = page.insert_textbox(rect, q_text, fontsize=10, fontname="helv", color=(0.1, 0.1, 0.1))
            # Rough line estimate
            lines_est = max(1, len(q_text) // 75 + q_text.count("\n"))
            y += max(18, lines_est * 13 + 6)

            # Options
            if opts:
                for opt in opts:
                    if y > cls.MARGIN_BOTTOM - 30:
                        page_num += 1
                        page = cls._create_page(doc, title, page_num)
                        y = cls.MARGIN_TOP

                    opt_key = opt.get("key", "")
                    opt_val = opt.get("text", "")
                    opt_str = f"({opt_key})  {opt_val}"
                    page.insert_text((cls.MARGIN_LEFT + 20, y + 10), opt_str, fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.2))
                    y += 14
                y += 6

            page.draw_line((cls.MARGIN_LEFT, y), (cls.MARGIN_RIGHT, y), color=(0.85, 0.85, 0.85), width=0.5)
            y += 10

        # Section 2: Answer Key & Detailed Solutions Guide
        if include_answers:
            page_num += 1
            page = cls._create_page(doc, title, page_num)
            y = cls.MARGIN_TOP

            page.insert_text((cls.MARGIN_LEFT, y + 14), "PART II: ANSWER KEY & DETAILED SOLUTIONS", fontsize=13, fontname="hebo", color=(0.1, 0.4, 0.2))
            y += 22
            page.draw_line((cls.MARGIN_LEFT, y), (cls.MARGIN_RIGHT, y), color=(0.1, 0.5, 0.25), width=1.5)
            y += 16

            # Quick Answer Grid Table
            page.insert_text((cls.MARGIN_LEFT, y + 10), "Quick Reference Key:", fontsize=10, fontname="hebo", color=(0.2, 0.2, 0.2))
            y += 16

            grid_cols = 5
            col_width = (cls.CONTENT_WIDTH) / grid_cols
            grid_y = y
            for idx, q in enumerate(questions):
                col_idx = idx % grid_cols
                row_idx = idx // grid_cols
                item_y = grid_y + row_idx * 16
                if item_y > cls.MARGIN_BOTTOM - 60:
                    page_num += 1
                    page = cls._create_page(doc, title, page_num)
                    grid_y = cls.MARGIN_TOP
                    item_y = grid_y + (idx % (grid_cols * 40)) // grid_cols * 16

                q_num = q.get("questionNumber") or str(idx + 1)
                ans = q.get("correctAnswer") or "-"
                cell_text = f"Q{q_num}: {ans}"
                cell_x = cls.MARGIN_LEFT + col_idx * col_width
                page.insert_text((cell_x, item_y + 10), cell_text, fontsize=9, fontname="hebo", color=(0.1, 0.3, 0.1))
                y = max(y, item_y + 18)

            y += 16
            page.draw_line((cls.MARGIN_LEFT, y), (cls.MARGIN_RIGHT, y), color=(0.8, 0.8, 0.8), width=0.5)
            y += 14

            # Detailed Explanations
            page.insert_text((cls.MARGIN_LEFT, y + 10), "Detailed Solutions & Explanations:", fontsize=10, fontname="hebo", color=(0.2, 0.2, 0.2))
            y += 16

            for idx, q in enumerate(questions):
                if y > cls.MARGIN_BOTTOM - 60:
                    page_num += 1
                    page = cls._create_page(doc, title, page_num)
                    y = cls.MARGIN_TOP

                q_num = q.get("questionNumber") or str(idx + 1)
                ans = q.get("correctAnswer") or "Not Specified"
                expl = q.get("explanation") or "No explanation provided."

                ans_header = f"Q{q_num}. Correct Answer: ({ans})" if len(ans) == 1 else f"Q{q_num}. Correct Answer: {ans}"
                page.insert_text((cls.MARGIN_LEFT, y + 10), ans_header, fontsize=9.5, fontname="hebo", color=(0.05, 0.35, 0.15))
                y += 15

                # Explanation text block
                rect = pymupdf.Rect(cls.MARGIN_LEFT + 15, y, cls.MARGIN_RIGHT, y + 100)
                page.insert_textbox(rect, expl, fontsize=9, fontname="helv", color=(0.25, 0.25, 0.25))
                lines_est = max(1, len(expl) // 80 + expl.count("\n"))
                y += max(16, lines_est * 12 + 8)

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    @classmethod
    def generate_paper_pdf(
        cls,
        paper_data: Dict[str, Any],
        include_answers: bool = False,
    ) -> bytes:
        """Generates an A4 examination question paper PDF with candidate header, questions, and optional answer key."""
        doc = pymupdf.open()
        page_num = 1
        title = paper_data.get("title", "Examination Paper")

        # Parse settings and compute margins in points (1 mm = 2.83465 pt)
        settings = paper_data.get("settings") or {}
        pt_per_mm = 2.83465
        preset = settings.get("pageMargin") or paper_data.get("pageMargin") or "normal"
        if preset == "zero":
            top_mm, right_mm, bottom_mm, left_mm = 4, 5, 4, 5
        elif preset == "narrow":
            top_mm, right_mm, bottom_mm, left_mm = 8, 10, 8, 10
        elif preset == "wide":
            top_mm, right_mm, bottom_mm, left_mm = 25, 25, 25, 25
        elif preset == "custom":
            top_mm = settings.get("marginTop", 15)
            right_mm = settings.get("marginRight", 18)
            bottom_mm = settings.get("marginBottom", 15)
            left_mm = settings.get("marginLeft", 18)
        else:
            top_mm, right_mm, bottom_mm, left_mm = 15, 18, 15, 18

        margin_left = int(left_mm * pt_per_mm)
        margin_right = cls.PAGE_WIDTH - int(right_mm * pt_per_mm)
        margin_top = int(top_mm * pt_per_mm)
        margin_bottom = cls.PAGE_HEIGHT - int(bottom_mm * pt_per_mm)

        page = cls._create_page(doc, title, page_num, margin_left, margin_right, margin_bottom)
        y = margin_top

        school_name = paper_data.get("schoolName") or "DELHI PUBLIC SCHOOL"
        exam_code = paper_data.get("examCode") or "EXAM-101"
        max_marks = paper_data.get("maxMarks", 100)
        duration = paper_data.get("durationMinutes", 180)
        exam_date = paper_data.get("examDate") or ""
        instructions = paper_data.get("instructions") or "1. Answer all questions.\n2. Write legibly."

        # Header Block
        page.insert_text((margin_left, y + 16), school_name.upper(), fontsize=14, fontname="hebo", color=(0.1, 0.1, 0.2))
        y += 22
        page.insert_text((margin_left, y + 12), title, fontsize=12, fontname="hebo", color=(0.15, 0.25, 0.5))
        y += 18

        # Metadata Table
        meta_left = f"EXAM CODE: {exam_code}   |   DATE: {exam_date}" if exam_date else f"EXAM CODE: {exam_code}"
        page.insert_text((margin_left, y + 10), meta_left, fontsize=9, fontname="hebo", color=(0.2, 0.2, 0.2))
        meta_right = f"TIME: {duration} MINS   |   MAX MARKS: {max_marks}"
        right_w = pymupdf.get_text_length(meta_right, fontname="hebo", fontsize=9)
        page.insert_text((margin_right - right_w, y + 10), meta_right, fontsize=9, fontname="hebo", color=(0.2, 0.2, 0.2))
        y += 16
        page.draw_line((margin_left, y), (margin_right, y), color=(0.1, 0.1, 0.1), width=1.5)
        y += 12

        # Candidate Details Box
        cand_rect = pymupdf.Rect(margin_left, y, margin_right, y + 42)
        page.draw_rect(cand_rect, color=(0.7, 0.7, 0.7), fill=(0.97, 0.97, 0.98), width=0.8)
        page.insert_text((margin_left + 8, y + 14), "Candidate Name: ________________________________________________", fontsize=9, fontname="helv")
        page.insert_text((margin_left + 8, y + 30), "Roll Number: [ ][ ][ ][ ][ ][ ][ ][ ]", fontsize=9, fontname="helv")
        page.insert_text((margin_right - 180, y + 30), "Signature: _______________________", fontsize=9, fontname="helv")
        y += 50

        # Instructions
        if instructions:
            page.insert_text((margin_left, y + 9), "General Instructions:", fontsize=8.5, fontname="hebo", color=(0.2, 0.2, 0.2))
            y += 13
            inst_rect = pymupdf.Rect(margin_left, y, margin_right, y + 40)
            page.insert_textbox(inst_rect, instructions, fontsize=8, fontname="helv", color=(0.3, 0.3, 0.3))
            y += 32
            page.draw_line((margin_left, y), (margin_right, y), color=(0.8, 0.8, 0.8), width=0.5)
            y += 12

        # Questions
        questions = paper_data.get("questions") or []
        for idx, q in enumerate(questions):
            if y > margin_bottom - 80:
                page_num += 1
                page = cls._create_page(doc, title, page_num, margin_left, margin_right, margin_bottom)
                y = margin_top

            q_num = q.get("questionNumber") or str(idx + 1)
            marks = q.get("marks", 1)
            q_text = q.get("questionText") or q.get("question_text") or ""
            opts = q.get("options") or []
            if isinstance(opts, str):
                try:
                    import json
                    opts = json.loads(opts)
                except Exception:
                    opts = []

            marks_str = f"[{marks}]"
            marks_w = pymupdf.get_text_length(marks_str, fontname="hebo", fontsize=9)
            page.insert_text((margin_right - marks_w, y + 10), marks_str, fontsize=9, fontname="hebo", color=(0.2, 0.2, 0.2))

            stem_prefix = f"Q{q_num}. "
            prefix_w = pymupdf.get_text_length(stem_prefix, fontname="hebo", fontsize=10)
            page.insert_text((margin_left, y + 10), stem_prefix, fontsize=10, fontname="hebo", color=(0.1, 0.1, 0.1))

            rect = pymupdf.Rect(margin_left + prefix_w, y, margin_right - marks_w - 8, y + 120)
            page.insert_textbox(rect, q_text, fontsize=9.5, fontname="helv", color=(0.1, 0.1, 0.1))
            lines_est = max(1, len(q_text) // 80 + q_text.count("\n"))
            y += max(18, lines_est * 13 + 6)

            # Options
            if opts:
                opt_strs = [f"({o.get('key', '')}) {o.get('text', '')}" for o in opts]
                for opt_s in opt_strs:
                    if y > margin_bottom - 25:
                        page_num += 1
                        page = cls._create_page(doc, title, page_num, margin_left, margin_right, margin_bottom)
                        y = margin_top
                    page.insert_text((margin_left + 20, y + 10), opt_s, fontsize=9, fontname="helv", color=(0.2, 0.2, 0.2))
                    y += 13
                y += 4

            page.draw_line((margin_left, y), (margin_right, y), color=(0.88, 0.88, 0.88), width=0.5)
            y += 8

        # Optional Answer Key Section
        if include_answers:
            page_num += 1
            page = cls._create_page(doc, f"{title} - Marking Scheme", page_num, margin_left, margin_right, margin_bottom)
            y = margin_top

            page.insert_text((margin_left, y + 14), "EXAMINATION ANSWER KEY & MARKING SCHEME", fontsize=13, fontname="hebo", color=(0.1, 0.4, 0.2))
            y += 22
            page.draw_line((margin_left, y), (margin_right, y), color=(0.1, 0.4, 0.2), width=1.5)
            y += 16

            for idx, q in enumerate(questions):
                if y > margin_bottom - 50:
                    page_num += 1
                    page = cls._create_page(doc, f"{title} - Marking Scheme", page_num, margin_left, margin_right, margin_bottom)
                    y = margin_top

                q_num = q.get("questionNumber") or str(idx + 1)
                ans = q.get("correctAnswer") or "-"
                expl = q.get("explanation") or "Full marks awarded for correct answer."
                marks = q.get("marks", 1)

                page.insert_text((margin_left, y + 10), f"Q{q_num}. Correct Answer: {ans}  [{marks} Mark{'s' if marks > 1 else ''}]", fontsize=9.5, fontname="hebo", color=(0.05, 0.35, 0.15))
                y += 14

                rect = pymupdf.Rect(margin_left + 15, y, margin_right, y + 60)
                page.insert_textbox(rect, f"Marking Scheme / Solution: {expl}", fontsize=8.5, fontname="helv", color=(0.3, 0.3, 0.3))
                y += 24

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    @classmethod
    def generate_text_pdf(cls, title: str, text: str, meta: Optional[Dict[str, Any]] = None) -> bytes:
        """Converts digital, extracted, or translated text into a formatted PDF document."""
        doc = pymupdf.open()
        page_num = 1
        page = cls._create_page(doc, title, page_num)
        y = cls.MARGIN_TOP

        # Title
        page.insert_text((cls.MARGIN_LEFT, y + 16), title.upper(), fontsize=14, fontname="hebo", color=(0.1, 0.15, 0.35))
        y += 24

        if meta:
            meta_parts = []
            if "language" in meta:
                meta_parts.append(f"Language: {meta['language']}")
            if "confidence" in meta:
                meta_parts.append(f"Confidence: {meta['confidence']}")
            if "lines_count" in meta:
                meta_parts.append(f"Lines: {meta['lines_count']}")
            if meta_parts:
                page.insert_text((cls.MARGIN_LEFT, y + 10), "  •  ".join(meta_parts), fontsize=8.5, fontname="helv", color=(0.4, 0.4, 0.4))
                y += 16

        page.draw_line((cls.MARGIN_LEFT, y), (cls.MARGIN_RIGHT, y), color=(0.2, 0.3, 0.6), width=1.5)
        y += 18

        lines = text.split("\n")
        for line in lines:
            line_str = line.rstrip()
            if not line_str:
                y += 8
                continue

            if y > cls.MARGIN_BOTTOM - 25:
                page_num += 1
                page = cls._create_page(doc, title, page_num)
                y = cls.MARGIN_TOP

            # Detect headings
            is_header = line_str.startswith("#") or (len(line_str) < 60 and line_str.isupper())
            font_args = get_unicode_font_args(is_bold=is_header)
            size = 11 if is_header else 9.5
            color = (0.1, 0.1, 0.3) if is_header else (0.1, 0.1, 0.1)

            clean_line = line_str.lstrip("#").strip()
            rect = pymupdf.Rect(cls.MARGIN_LEFT, y, cls.MARGIN_RIGHT, y + 40)
            page.insert_textbox(rect, clean_line, fontsize=size, color=color, **font_args)
            lines_est = max(1, len(clean_line) // 85)
            y += lines_est * (size * 1.35) + 3

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    @classmethod
    def parse_questions_from_text(cls, text: str) -> List[Dict[str, Any]]:
        """
        Parses raw text into structured questions and answers:
        - Question number & stem
        - Options (A), (B), (C), (D)
        - Correct answer: 'Ans: A', 'Answer: (B)', 'Correct: C'
        - Explanation / Solution: 'Explanation: ...', 'Solution: ...'
        - Marks: '[2 Marks]', '(3 pts)'
        """
        if not text or not text.strip():
            return []

        lines = text.split("\n")
        questions: List[Dict[str, Any]] = []
        current_q: Optional[Dict[str, Any]] = None

        q_regex = re.compile(r"^(?:Q(?:uestion)?\.?\s*(\d+)(?:[.)\]:\-\s]|\b)|\((\d+)\)|(\d+)[.)\]:\-])\s*(.*)$", re.IGNORECASE)
        opt_inline_regex = re.compile(r"(?:\(([a-dA-D1-4])\)|(?:(?<=^)|(?<=\s))([a-dA-D1-4])\.(?!\d)|(?:(?<=^)|(?<=\s))([a-dA-D1-4])\))\s*")
        ans_regex = re.compile(r"(?:Ans(?:wer)?|Correct\s*(?:Option)?|Key)[:\s]+(?:\(?([A-Da-d])\)?|(.*))", re.IGNORECASE)
        expl_regex = re.compile(r"^(?:Explanation|Solution|Hint|Reason)[:\s]+(.*)$", re.IGNORECASE)
        marks_regex = re.compile(r"(?:\[|\()(\d+)\s*(?:marks?|mark|m|pts?)(?:\]|\))", re.IGNORECASE)

        for line in lines:
            trimmed = line.strip()
            if not trimmed:
                continue

            # Check for new question start
            q_match = q_regex.match(trimmed)
            if q_match:
                if current_q:
                    questions.append(current_q)
                
                q_num = q_match.group(1) or q_match.group(2) or q_match.group(3) or str(len(questions) + 1)
                stem_text = q_match.group(4) or ""
                
                # Check marks in stem
                marks = 1
                m_match = marks_regex.search(stem_text)
                if m_match:
                    marks = int(m_match.group(1))
                    stem_text = marks_regex.sub("", stem_text).strip()

                # Check if inline options are concatenated at the end of stem
                inline_opts = []
                opt_matches = list(opt_inline_regex.finditer(stem_text))
                if len(opt_matches) >= 2:
                    first_opt_idx = opt_matches[0].start()
                    opts_text = stem_text[first_opt_idx:]
                    stem_text = stem_text[:first_opt_idx].strip()
                    m_list = list(opt_inline_regex.finditer(opts_text))
                    for i, om in enumerate(m_list):
                        label = (om.group(1) or om.group(2) or om.group(3)).upper()
                        mapping = {"1": "A", "2": "B", "3": "C", "4": "D"}
                        label = mapping.get(label, label)
                        start = om.end()
                        end = m_list[i + 1].start() if i + 1 < len(m_list) else len(opts_text)
                        val = opts_text[start:end].strip()
                        inline_opts.append({
                            "key": label,
                            "text": specialized_math.convert_embedded_math(val)
                        })

                normalized_stem = specialized_math.convert_embedded_math(stem_text)

                current_q = {
                    "questionNumber": q_num,
                    "questionText": normalized_stem,
                    "options": inline_opts,
                    "correctAnswer": "",
                    "explanation": "",
                    "marks": marks,
                    "difficulty": "MEDIUM",
                }
                continue

            if not current_q:
                # If header/title before first question, ignore header instead of treating it as question 1
                is_paper_header = (
                    re.search(r"^(?:DPP|DAILY|EXAMINATION|EXAM|TEST|PAPER|CLASS|SUBJECT|TIME|MAX|MARKS|SECTION|PART|INSTRUCTION|ANNUAL|MID-TERM|HALF|MATHEMATICS|SCIENCE|PHYSICS|CHEMISTRY|BIOLOGY)", trimmed, re.IGNORECASE)
                    or (trimmed.isupper() and len(trimmed) < 80)
                )
                if is_paper_header:
                    continue

                # Start default question 1 only if actual question content
                current_q = {
                    "questionNumber": "1",
                    "questionText": specialized_math.convert_embedded_math(trimmed),
                    "options": [],
                    "correctAnswer": "",
                    "explanation": "",
                    "marks": 1,
                    "difficulty": "MEDIUM",
                }
                continue

            # Check for Explanation
            expl_match = expl_regex.match(trimmed)
            if expl_match:
                current_q["explanation"] = specialized_math.convert_embedded_math(expl_match.group(1).strip())
                continue

            # Check for Correct Answer
            ans_match = ans_regex.search(trimmed)
            if ans_match:
                ans_key = (ans_match.group(1) or ans_match.group(2) or "").strip().upper()
                current_q["correctAnswer"] = ans_key
                # Also check marks in answer line
                m_match = marks_regex.search(trimmed)
                if m_match:
                    current_q["marks"] = int(m_match.group(1))
                continue

            # Check for Options on separate line (could be 1 or multiple inline e.g. (a) ... (b) ...)
            opt_matches = list(opt_inline_regex.finditer(trimmed))
            if opt_matches:
                for i, om in enumerate(opt_matches):
                    label = (om.group(1) or om.group(2) or om.group(3)).upper()
                    mapping = {"1": "A", "2": "B", "3": "C", "4": "D"}
                    label = mapping.get(label, label)
                    start = om.end()
                    end = opt_matches[i + 1].start() if i + 1 < len(opt_matches) else len(trimmed)
                    val = trimmed[start:end].strip()
                    current_q["options"].append({
                        "key": label,
                        "text": specialized_math.convert_embedded_math(val)
                    })
                continue

            # Otherwise, append to existing question text or explanation
            if current_q["explanation"]:
                current_q["explanation"] += "\n" + specialized_math.convert_embedded_math(trimmed)
            elif not current_q["options"]:
                current_q["questionText"] += "\n" + specialized_math.convert_embedded_math(trimmed)
            else:
                # Append to last option text
                current_q["options"][-1]["text"] += " " + specialized_math.convert_embedded_math(trimmed)

        if current_q:
            questions.append(current_q)

        return questions

    @classmethod
    def parse_word_file(cls, docx_path: Path) -> List[Dict[str, Any]]:
        """Extracts text and tables from Word .docx document preserving OMML math equations and parses questions and answers."""
        try:
            doc = docx.Document(docx_path)
            full_text_lines: List[str] = []

            for p in doc.paragraphs:
                text = extract_paragraph_with_math(p)
                if text:
                    full_text_lines.append(text)

            for table in doc.tables:
                for row in table.rows:
                    row_texts = [extract_cell_with_math(cell) for cell in row.cells if cell.text.strip() or len(cell.paragraphs) > 0]
                    row_texts = [rt for rt in row_texts if rt.strip()]
                    if row_texts:
                        full_text_lines.append("  |  ".join(row_texts))

            full_text = "\n".join(full_text_lines)
            return cls.parse_questions_from_text(full_text)
        except Exception as e:
            logger.error(f"Error parsing Word file {docx_path}: {e}")
            raise

    @classmethod
    def parse_pdf_file(cls, pdf_path: Path) -> List[Dict[str, Any]]:
        """Extracts text from PDF document and parses questions and answers."""
        try:
            doc = pymupdf.open(pdf_path)
            full_text_lines: List[str] = []

            for page in doc:
                text = page.get_text("text")
                if text:
                    full_text_lines.append(text.strip())

            doc.close()
            full_text = "\n".join(full_text_lines)
            return cls.parse_questions_from_text(full_text)
        except Exception as e:
            logger.error(f"Error parsing PDF file {pdf_path}: {e}")
            raise

    @classmethod
    def extract_file_raw_text(cls, file_path: Path) -> str:
        """Extracts complete text from Word (.docx), PDF (.pdf), or plain text files with math formulas preserved."""
        try:
            ext = file_path.suffix.lower()
            if ext in [".docx", ".doc"]:
                doc = docx.Document(file_path)
                lines: List[str] = []
                for p in doc.paragraphs:
                    text = extract_paragraph_with_math(p)
                    if text:
                        lines.append(text)
                for table in doc.tables:
                    for row in table.rows:
                        row_texts = [extract_cell_with_math(cell) for cell in row.cells if cell.text.strip() or len(cell.paragraphs) > 0]
                        row_texts = [rt for rt in row_texts if rt.strip()]
                        if row_texts:
                            lines.append("  |  ".join(row_texts))
                return "\n".join(lines)
            elif ext == ".pdf":
                doc = pymupdf.open(file_path)
                lines = []
                for page in doc:
                    text = page.get_text("text")
                    if text:
                        lines.append(text.strip())
                doc.close()
                return "\n".join(lines)
            else:
                return file_path.read_text(encoding="utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"Error extracting raw text from {file_path}: {e}")
            raise

    @classmethod
    def generate_questions_csv(cls, questions: List[Dict[str, Any]]) -> str:
        """
        Exports questions to CSV according to Section 34 specifications:
        Question, Option A, Option B, Option C, Option D, Answer, Marks,
        Formula LaTeX, Formula Plain Text, Physics Formula, Chemistry Formula
        """
        import csv
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow([
            "Question",
            "Option A",
            "Option B",
            "Option C",
            "Option D",
            "Answer",
            "Marks",
            "Formula LaTeX",
            "Formula Plain Text",
            "Physics Formula",
            "Chemistry Formula",
        ])

        for idx, q in enumerate(questions):
            q_text = q.get("questionText") or q.get("question_text") or ""
            opts = q.get("options") or []
            if isinstance(opts, str):
                try:
                    import json
                    opts = json.loads(opts)
                except Exception:
                    opts = []

            opt_map: Dict[str, str] = {}
            for o in opts:
                k = (o.get("key") or "").upper().strip()
                t = o.get("text") or ""
                opt_map[k] = t

            ans = q.get("correctAnswer") or q.get("answer") or ""
            marks = q.get("marks", 1)

            # Extract formulas from question text and options
            full_content = q_text + " " + " ".join(opt_map.values())
            latex_formulas = re.findall(r"\$(.*?)\$", full_content)
            formula_latex = "; ".join(latex_formulas) if latex_formulas else ""
            formula_plain = "; ".join([f.replace("\\sqrt", "√").replace("\\times", "×").replace("\\div", "÷") for f in latex_formulas]) if latex_formulas else ""

            # Classify physics/chemistry
            phys_formula = ""
            chem_formula = ""
            if any(term in full_content for term in ["F = ma", "v = u", "E = mc", "m/s", "kg", "N", "J", "W", "Ω", "θ"]):
                phys_formula = formula_latex or full_content
            if any(term in full_content for term in ["H2O", "CO2", "H2SO4", "->", "→", "⇌", "Fe^", "SO4", "CH3COOH"]):
                chem_formula = formula_latex or full_content

            writer.writerow([
                q_text,
                opt_map.get("A", ""),
                opt_map.get("B", ""),
                opt_map.get("C", ""),
                opt_map.get("D", ""),
                ans,
                marks,
                formula_latex,
                formula_plain,
                phys_formula,
                chem_formula,
            ])

        return output.getvalue()

